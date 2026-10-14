#!/usr/bin/env python3
"""Gera docs/Documentacao_DevOps.pptx (documentação técnica da entrega DevOps).

Os prints em docs/evidencias/<arquivo>.png entram automaticamente no slide certo;
enquanto um print não existir, o slide mostra um espaço identificado para colá-lo.

    pip install python-pptx pillow
    python3 scripts/gerar_documentacao.py
"""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parent.parent
EVIDENCIAS = ROOT / "docs" / "evidencias"
SAIDA = ROOT / "docs" / "Documentacao_DevOps.pptx"

INTEGRANTES = "[PREENCHA COM OS NOMES E RMs]"
REPO_URL = "https://github.com/vsilva1997/orbital-sos-api"
STAGING_URL = "https://orbital-sos-api-staging.onrender.com"
PROD_URL = "https://orbital-sos-api-prod.onrender.com"
# Status do checklist: True = concluído
DEPLOY_OK = False

# Paleta: teal profundo (estabilidade/ambiente) + laranja SOS (alerta)
DARK, DARK2, TINT, ACCENT = "0E3B43", "155E63", "E8F3F1", "E8572A"
TEXT, MUTED, LIGHT_TXT, OK_GREEN = "1E2B2F", "56686C", "CFE5E1", "2E7D4F"
FONT, CODE_FONT = "Arial", "Courier New"
W, H = 13.333, 7.5


def rgb(hex_):
    return RGBColor.from_string(hex_)


# ---------- texto ----------
def _apply_bullet(p, color=ACCENT):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(0.28)))
    pPr.set("indent", str(-Inches(0.22)))
    clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(clr, qn("a:srgbClr")).set("val", color)
    etree.SubElement(pPr, qn("a:buFont")).set("typeface", FONT)
    etree.SubElement(pPr, qn("a:buChar")).set("char", "•")


def fill_text_frame(tf, paras, size=16, bold=False, color=TEXT, align=PP_ALIGN.LEFT, font=FONT,
                    bullets=False, space_after=6):
    """paras: str | lista de parágrafos; parágrafo = str | lista de runs; run = str | (texto, opções)."""
    paras = paras if isinstance(paras, list) else [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        runs = [para] if isinstance(para, (str, tuple)) else para
        for run in runs:
            txt, opt = (run, {}) if isinstance(run, str) else run
            r = p.add_run()
            r.text = txt
            f = r.font
            f.name = opt.get("font", font)
            f.size = Pt(opt.get("size", size))
            f.bold = opt.get("bold", bold)
            f.italic = opt.get("italic", False)
            f.color.rgb = rgb(opt.get("color", color))
        if bullets:
            _apply_bullet(p)


def text(slide, x, y, w, h, paras, size=16, bold=False, color=TEXT, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=FONT, bullets=False, space_after=6, margin=0.0, name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = anchor
    fill_text_frame(tf, paras, size, bold, color, align, font, bullets, space_after)
    return tb


# ---------- formas ----------
def box(slide, x, y, w, h, fill=TINT, line=None, radius=0.12, dash=False, line_w=1.25, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = min(0.5, radius / min(w, h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(line_w)
        if dash:
            shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    shp.shadow.inherit = False
    return shp


def node(slide, x, y, w, h, title, sub=None, fill=TINT, color=TEXT, sub_color=None, size=15, sub_size=12):
    """Caixa com título e subtítulo centralizados (diagramas)."""
    shp = box(slide, x, y, w, h, fill=fill)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.08)
    paras = [[(title, {"bold": True, "size": size})]]
    for line in (sub.split("\n") if sub else []):
        paras.append([(line, {"size": sub_size, "color": sub_color or color})])
    fill_text_frame(tf, paras, color=color, align=PP_ALIGN.CENTER, space_after=2)
    return shp


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=2.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    tail = etree.SubElement(c.line._get_or_add_ln(), qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return c


def numbered_circle(slide, x, y, n, d=0.48, fill=ACCENT):
    c = box(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    fill_text_frame(tf, [[(str(n), {"bold": True, "size": 16, "color": "FFFFFF"})]], align=PP_ALIGN.CENTER, space_after=0)
    return c


def table(slide, x, y, col_w, rows, size=14, header_fill=DARK, row_h=0.5, first_col_bold=True):
    shape = slide.shapes.add_table(len(rows), len(col_w), Inches(x), Inches(y), Inches(sum(col_w)), Inches(row_h * len(rows)))
    tbl = shape.table
    tbl.first_row = True
    tbl.horz_banding = False
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(header_fill if i == 0 else ("FFFFFF" if i % 2 else TINT))
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            if isinstance(val, tuple):  # (texto, cor) para status
                val, clr, bld = val[0], val[1], True
            else:
                clr = "FFFFFF" if i == 0 else TEXT
                bld = i == 0 or (j == 0 and first_col_bold)
            fill_text_frame(tf, [[(val, {"bold": bld, "color": clr})]], size=size, space_after=0)
    return shape


def evidence(slide, x, y, w, h, fname, caption):
    """Insere docs/evidencias/<fname> ajustado à área, ou um espaço tracejado para colar o print."""
    img_h = h - 0.45
    path = EVIDENCIAS / fname
    if path.exists():
        iw, ih = Image.open(path).size
        scale = min(w / iw, img_h / ih)
        pw, ph = iw * scale, ih * scale
        pic = slide.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y + (img_h - ph) / 2), Inches(pw), Inches(ph))
        pic.line.color.rgb = rgb("C9D9D6")
        pic.line.width = Pt(0.75)
    else:
        b = box(slide, x, y, w, img_h, fill="F5F8F8", line="9AAEB1", dash=True, radius=0.1)
        b.name = f"Espaço para {fname}"
        tf = b.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        fill_text_frame(tf, [[("Cole aqui o print", {"bold": True, "size": 16, "color": MUTED})],
                             [(fname, {"size": 14, "color": ACCENT, "font": CODE_FONT})]],
                        align=PP_ALIGN.CENTER, space_after=4)
    text(slide, x, y + img_h + 0.1, w, 0.35, caption, size=12, color=MUTED, align=PP_ALIGN.CENTER)


# ---------- molduras ----------
class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.blank = self.prs.slide_layouts[6]
        self.content_slides = []

    def dark(self):
        s = self.prs.slides.add_slide(self.blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(DARK)
        return s

    def content(self, kicker, title):
        s = self.prs.slides.add_slide(self.blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb("FFFFFF")
        text(s, 0.6, 0.38, 12.1, 0.3, kicker.upper(), size=12, bold=True, color=ACCENT, name="Seção")
        text(s, 0.6, 0.68, 12.1, 0.7, title, size=30, bold=True, color=DARK, name="Título")
        self.content_slides.append(s)
        return s

    def footers(self):
        total = len(self.prs.slides)
        for idx, s in enumerate(self.prs.slides, start=1):
            if s in self.content_slides:
                text(s, 0.6, 7.0, 8, 0.3, "Orbital SOS API · Cidades ESGInteligentes · FIAP", size=10, color=MUTED)
                text(s, 10.7, 7.0, 2.0, 0.3, f"{idx} / {total}", size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def build():
    d = Deck()

    # 1. Capa
    s = d.dark()
    text(s, 0.8, 1.15, 8.2, 0.4, "FIAP · NAVEGANDO PELO MUNDO DEVOPS", size=14, bold=True, color=ACCENT)
    text(s, 0.8, 1.6, 8.2, 1.0, "Cidades ESGInteligentes", size=46, bold=True, color="FFFFFF")
    text(s, 0.8, 2.7, 7.9, 1.1, "Orbital SOS API: containerização, pipeline CI/CD e deploy em staging e produção",
         size=22, color=LIGHT_TXT)
    card = box(s, 0.8, 4.35, 7.6, 1.5, fill=DARK2)
    card.name = "Integrantes"
    text(s, 1.1, 4.55, 7.0, 0.35, "INTEGRANTES", size=12, bold=True, color=ACCENT)
    text(s, 1.1, 4.95, 7.0, 0.7, INTEGRANTES, size=20, bold=True, color="FFFFFF")
    text(s, 0.8, 6.45, 8.0, 0.35, f"Outubro de 2026 · {REPO_URL.replace('https://', '')}", size=12, color=LIGHT_TXT)
    ring = box(s, 9.35, 1.55, 3.3, 3.3, fill=None, line=DARK2, line_w=4, shape=MSO_SHAPE.OVAL)
    ring.name = "Órbita"
    box(s, 12.2, 2.05, 0.5, 0.5, fill=ACCENT, shape=MSO_SHAPE.OVAL).name = "Satélite"
    core = box(s, 10.25, 2.45, 1.5, 1.5, fill="FFFFFF", shape=MSO_SHAPE.OVAL)
    core.name = "SOS"
    tf = core.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    fill_text_frame(tf, [[("SOS", {"bold": True, "size": 28, "color": DARK})]], align=PP_ALIGN.CENTER, space_after=0)

    # 2. Visão geral
    s = d.content("Contexto", "Conectando vítimas de desastres climáticos à ajuda")
    text(s, 0.6, 1.65, 6.5, 1.3,
         "A Orbital SOS API é o componente social do projeto Cidades ESGInteligentes: cadastra abrigos, alertas "
         "climáticos, áreas de risco, pedidos de socorro (SOS) e doações.", size=17)
    text(s, 0.6, 3.05, 6.5, 3.3, [
        ".NET 8 em Clean Architecture: API, Application, Domain e Infrastructure",
        "EF Core 8 com PostgreSQL (Supabase), autenticação JWT e Swagger",
        "Objetivo desta entrega: levar a API do código ao ar com Docker, pipeline CI/CD e dois ambientes na nuvem",
        "Tudo versionado no GitHub: código, Dockerfile, compose, pipeline e infraestrutura (render.yaml)",
    ], size=16, bullets=True, space_after=10)
    stats = [("14", "testes automatizados rodando em todo pull request"),
             ("5", "jobs no pipeline do GitHub Actions"),
             ("2", "ambientes na nuvem: staging e produção")]
    for i, (num, label) in enumerate(stats):
        y = 1.65 + i * 1.65
        box(s, 7.6, y, 5.1, 1.35, fill=TINT).name = f"Indicador {num}"
        text(s, 7.85, y + 0.17, 1.3, 1.0, num, size=44, bold=True, color=ACCENT, anchor=MSO_ANCHOR.MIDDLE)
        text(s, 9.15, y + 0.17, 3.35, 1.0, label, size=16, color=TEXT, anchor=MSO_ANCHOR.MIDDLE)

    # 3. Fluxo do pipeline
    s = d.content("Pipeline CI/CD", "Do pull request à produção com GitHub Actions")
    nw, gap = 1.7, 0.38
    xs = [0.6 + i * (nw + gap) for i in range(6)]
    rows = [
        ("Staging: branch develop, deploy automático", 2.1, [
            (0, "feature/devops", "branch", DARK, "FFFFFF"),
            (1, "Pull request", "build · testes · compose", TINT, TEXT),
            (2, "develop", "merge", DARK, "FFFFFF"),
            (3, "Imagem :sha", "push no GHCR", TINT, TEXT),
            (5, "Deploy staging", "smoke /health", DARK2, "FFFFFF")]),
        ("Produção: branch main, com aprovação manual", 4.45, [
            (0, "develop", "branch", DARK, "FFFFFF"),
            (1, "Pull request", "build · testes · compose", TINT, TEXT),
            (2, "main", "merge", DARK, "FFFFFF"),
            (3, "Imagem :sha", "push no GHCR", TINT, TEXT),
            (4, "Aprovação", "environment production", ACCENT, "FFFFFF"),
            (5, "Deploy produção", "smoke /health", DARK2, "FFFFFF")]),
    ]
    nh = 1.1
    for label, y, nodes in rows:
        text(s, 0.6, y - 0.45, 12.1, 0.35, label, size=15, bold=True, color=MUTED)
        for k, (pos, t, sub, fill, clr) in enumerate(nodes):
            node(s, xs[pos], y, nw, nh, t, sub, fill=fill, color=clr, size=15, sub_size=12)
            if k + 1 < len(nodes):
                nxt = nodes[k + 1][0]
                arrow(s, xs[pos] + nw + 0.04, y + nh / 2, xs[nxt] - 0.04, y + nh / 2)
        if label.startswith("Staging"):
            text(s, xs[4], y + 0.05, nw, 0.4, "automático", size=12, color=MUTED, align=PP_ALIGN.CENTER)
    text(s, 0.6, 6.25, 12.1, 0.4, [[("Ferramenta: ", {"bold": True}), "GitHub Actions   ·   ",
                                    ("Registro: ", {"bold": True}), "GitHub Container Registry   ·   ",
                                    ("Hospedagem: ", {"bold": True}), "Render + Supabase (PostgreSQL)"]],
         size=15, color=TEXT)

    # 4. Etapas do pipeline
    s = d.content("Pipeline CI/CD", "Cinco jobs, cada um com uma responsabilidade")
    table(s, 0.6, 1.6, [2.35, 2.55, 7.2], [
        ["Job", "Quando roda", "O que faz"],
        ["build-test", "PR e push em develop/main", "Restore com lock file e cache do NuGet, build Release, testes xUnit (resultado .trx como artefato) e checagem de pacotes vulneráveis"],
        ["compose-smoke", "PR e push", "Valida os 3 arquivos compose, sobe API + PostgreSQL 17 e testa health, swagger, register, login e admin"],
        ["docker", "Só em push", "Build multi-stage com cache e push de ghcr.io/vsilva1997/orbital-sos-api:<sha>"],
        ["deploy-staging", "Push em develop", "Deploy hook do Render com a imagem do SHA; espera o /health responder Healthy na versão nova"],
        ["deploy-production", "Push em main", "Mesmo deploy, depois da aprovação manual no environment production"],
    ], size=14, row_h=0.55)
    callout = box(s, 0.6, 5.95, 12.1, 0.8, fill=TINT)
    callout.name = "Lógica do pipeline"
    text(s, 0.85, 6.03, 11.6, 0.65, [[("Lógica: ", {"bold": True, "color": ACCENT}),
                                      "o PR só valida. Depois do merge, a imagem é gerada uma vez por commit (tag = SHA) e o /health "
                                      "devolve essa versão: o smoke test só passa quando o deploy certo está no ar."]],
         size=14, anchor=MSO_ANCHOR.MIDDLE)

    # 5. Arquitetura Docker
    s = d.content("Docker", "A mesma imagem roda no compose local e na nuvem")
    text(s, 0.6, 1.5, 6.0, 0.35, "Local: docker compose", size=15, bold=True, color=MUTED)
    net = box(s, 0.6, 1.95, 6.0, 4.75, fill=None, line="8FB3AE", dash=True, radius=0.15)
    net.name = "Rede orbital-net"
    text(s, 0.85, 2.05, 5.5, 0.3, "rede orbital-net (bridge)", size=12, color=MUTED, font=CODE_FONT)
    node(s, 0.85, 2.5, 2.45, 1.55, "api", "OrbitalSOS.API\nporta 8080 · usuário app\nHEALTHCHECK /health",
         fill=DARK, color="FFFFFF", sub_color=LIGHT_TXT, size=17, sub_size=12)
    node(s, 3.9, 2.5, 2.45, 1.55, "db", "PostgreSQL 17\nhealthcheck pg_isready\nporta só em 127.0.0.1",
         fill=DARK2, color="FFFFFF", sub_color=LIGHT_TXT, size=17, sub_size=12)
    arrow(s, 3.33, 3.27, 3.87, 3.27)
    text(s, 2.3, 4.2, 2.6, 0.55, ["depends_on: db", "service_healthy"], size=12, color=MUTED, align=PP_ALIGN.CENTER, space_after=0)
    host = box(s, 0.85, 4.85, 2.45, 1.05, fill=TINT)
    host.name = "Host"
    text(s, 0.95, 4.92, 2.25, 0.9, [[("localhost:8080", {"bold": True})], "Swagger e /health"], size=13,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=2)
    arrow(s, 2.075, 4.82, 2.075, 4.08)
    vol = box(s, 4.2, 4.75, 1.85, 1.25, fill=TINT, line="8FB3AE", shape=MSO_SHAPE.CAN)
    vol.name = "Volume pg-data"
    text(s, 4.05, 5.12, 2.15, 0.8, [[("pg-data", {"bold": True})], "/var/lib/postgresql"], size=12,
         align=PP_ALIGN.CENTER, space_after=0)
    arrow(s, 5.125, 4.08, 5.125, 4.72)

    text(s, 7.1, 1.5, 5.6, 0.35, "Nuvem: staging e produção", size=15, bold=True, color=MUTED)
    node(s, 7.1, 1.95, 5.6, 1.0, "ghcr.io/vsilva1997/orbital-sos-api", "tags :<sha> · :develop · :main",
         fill=DARK, color="FFFFFF", sub_color=LIGHT_TXT, size=15, sub_size=12)
    for x, env, schema, user in ((7.1, "Staging", "staging", "orbital_staging"), (10.05, "Production", "production", "orbital_prod")):
        node(s, x, 3.5, 2.65, 1.2, "Render", f"Web Service Free\nambiente {env}", fill=TINT, size=15, sub_size=11)
        node(s, x, 5.3, 2.65, 1.1, "Supabase", f"schema {schema}\nusuário {user}", fill=TINT, size=15, sub_size=12)
        arrow(s, x + 1.325, 2.98, x + 1.325, 3.47)
        arrow(s, x + 1.325, 4.73, x + 1.325, 5.27)

    # 6. Dockerfile, imagem e comandos
    s = d.content("Docker", "Imagem enxuta, sem root e com health check")
    decisions = [
        ("Multi-stage sdk:8.0 → aspnet:8.0", "O SDK fica só no build; a imagem final leva apenas o runtime"),
        ("Restore em camada separada", "csproj e lock files primeiro: mudar código não baixa os pacotes de novo"),
        ("Usuário não root e porta 8080", "Padrão do .NET 8 (USER app); antes rodava como root na porta 80"),
        ("HEALTHCHECK em /health", "curl instalado só para isso; o /health testa a API e o banco"),
        ("APP_VERSION = SHA do commit", "Aparece no /health e o smoke test confere a versão no ar"),
    ]
    for i, (t, desc) in enumerate(decisions):
        y = 1.6 + i * 1.0
        numbered_circle(s, 0.6, y + 0.05, i + 1)
        text(s, 1.25, y, 5.6, 0.95, [[(t, {"bold": True, "size": 16})], [(desc, {"size": 14, "color": MUTED})]], space_after=2)
    box(s, 7.3, 1.6, 5.4, 1.4, fill=TINT).name = "Imagem criada"
    text(s, 7.5, 1.7, 5.0, 1.25, [[("Imagem criada", {"bold": True, "size": 14, "color": ACCENT})],
                                  [("ghcr.io/vsilva1997/orbital-sos-api", {"font": CODE_FONT, "size": 13, "bold": True})],
                                  [(":<sha> (deploy) · :develop · :main", {"size": 14, "color": MUTED})]], space_after=4)
    code = box(s, 7.3, 3.25, 5.4, 3.35, fill=DARK, radius=0.1)
    code.name = "Comandos"
    text(s, 7.55, 3.4, 5.0, 3.1, [
        [("# comandos usados", {"color": "8FC1B9"})],
        "docker build -t orbital-sos-api:local .",
        "docker compose up -d --build",
        "docker compose ps",
        "docker compose logs -f api",
        "docker compose down -v",
        "docker pull ghcr.io/vsilva1997/",
        "    orbital-sos-api:main",
        [("# simular staging localmente", {"color": "8FC1B9"})],
        "docker compose -f docker-compose.yml \\",
        "  -f docker-compose.staging.yml up -d",
    ], size=13, color="FFFFFF", font=CODE_FONT, space_after=3)

    # 7-13. Evidências
    pares = [
        ("Evidências: pipeline", "Build e testes validados em todo pull request",
         ("01-pr-checks.png", "PR feature/devops → develop com os checks verdes"),
         ("03-testes-resumo.png", "Resumo do job: 14/14 testes aprovados e artefato .trx")),
        ("Evidências: pipeline", "Compose testado no CI e imagem publicada no GHCR",
         ("04-compose-smoke.png", "Job compose-smoke: containers healthy e smoke test OK"),
         ("05-ghcr-imagem.png", "Pacote no GHCR com as tags do SHA, develop e main")),
        ("Evidências: deploy", "Cada merge publica no ambiente da sua branch",
         ("02-pipeline-develop.png", "Run da develop até o deploy em staging"),
         ("07-pipeline-main.png", "Run da main até o deploy em produção")),
        ("Evidências: deploy", "Produção só recebe deploy após aprovação manual",
         ("06-aprovacao-producao.png", "Review deployments no environment production"),
         ("08-environments.png", "Environments com regras de branch e revisor obrigatório")),
        ("Evidências: staging", "Staging no ar com a versão do último merge na develop",
         ("09-staging-health.png", "/health de staging: Healthy, Staging e SHA do commit"),
         ("10-staging-swagger.png", "Swagger de staging")),
        ("Evidências: produção", "Produção no ar, com register e login funcionando",
         ("11-producao-health.png", "/health de produção: Healthy e Production"),
         ("12-producao-swagger.png", "Register e login executados no Swagger de produção")),
        ("Evidências: nuvem", "Infraestrutura dos dois ambientes",
         ("13-render-servicos.png", "Render: serviços de staging e produção no ar"),
         ("14-supabase.png", "Supabase: schemas de staging e produção (plano Free)")),
    ]
    for kicker, title, left, right in pares:
        s = d.content(kicker, title)
        evidence(s, 0.6, 1.6, 5.9, 5.15, *left)
        evidence(s, 6.8, 1.6, 5.9, 5.15, *right)

    # 14-15. Desafios
    desafios = [
        ("Desafios e soluções", "Código: o projeto precisou de ajustes antes do pipeline", [
            ["Desafio encontrado", "Como foi resolvido"],
            ["O projeto não compilava: faltava ImplicitUsings e a referência Application ↔ Infrastructure estava invertida",
             "Habilitar ImplicitUsings e inverter a referência, sem mudar regra de negócio"],
            ["Os 8 testes falhavam: Database.Migrate() não funciona com EF InMemory",
             "Migrar só quando Database.IsRelational(); banco InMemory estável por factory"],
            ["Senha do SA e chave JWT fixas no Git",
             "Variáveis de ambiente (.env fora do Git, secrets no Render e no GitHub); valores antigos descartados"],
            ["Qualquer pessoa podia se cadastrar como Admin (campo role no register)",
             "Register sempre cria Usuario; Admin inicial via Admin__Email e Admin__Password"],
            ["Pacotes com vulnerabilidades High (SqlClient, System.Text.Json)",
             "Atualização para o patch 8.0.32 e checagem automática no pipeline"],
        ]),
        ("Desafios e soluções", "Infraestrutura: containers, CI e nuvem gratuita", [
            ["Desafio encontrado", "Como foi resolvido"],
            ["O Azure for Students não pôde ser ativado, e o SQL Server (2 GB, x86) não tem hospedagem gratuita fora dele",
             "Troca para PostgreSQL: provider Npgsql, migration regerada, compose com postgres:17 e banco no Supabase Free"],
            ["Container rodando como root na porta 80", "USER app e porta 8080, padrão do .NET 8"],
            ["Mac sem Docker instalado para testar o compose",
             "Compose validado no runner do GitHub Actions em todo PR (job compose-smoke)"],
            ["Plano Free do Supabase permite só 2 projetos ativos",
             "Um projeto com schemas staging e production, cada um com usuário próprio e senha em hash SCRAM"],
            ["Cold start: o Render dorme e o Supabase pausa após 7 dias sem uso",
             "Retry no EF, smoke test com retry e health check diário agendado no GitHub Actions"],
            ["ubuntu-latest migraria para Ubuntu 26 na semana da entrega", "Runner fixado em ubuntu-24.04"],
        ]),
    ]
    for kicker, title, rows in desafios:
        s = d.content(kicker, title)
        table(s, 0.6, 1.6, [6.05, 6.05], rows, size=14, row_h=0.62, first_col_bold=False)

    # 16. Checklist
    s = d.content("Entrega", "Checklist de entrega")
    ok, pend = ("Concluído", OK_GREEN), ("Pendente", ACCENT)
    tem_prints = all((EVIDENCIAS / f).exists() for _, _, esq, dir_ in pares for f in (esq[0], dir_[0]))
    table(s, 0.6, 1.6, [6.4, 1.7, 4.0], [
        ["Item", "Status", "Onde"],
        ["Projeto compactado em .ZIP com estrutura organizada", ok, "scripts/empacotar.sh → entrega.zip"],
        ["Dockerfile funcional", ok, "Dockerfile (validado no pipeline)"],
        ["docker-compose.yml ou arquivos Kubernetes", ok, "docker-compose*.yml"],
        ["Pipeline com etapas de build, teste e deploy", ok, ".github/workflows/ci-cd.yml"],
        ["README.md com instruções e prints", ok if tem_prints else pend, "README.md + docs/evidencias/"],
        ["Documentação técnica com evidências (PDF ou PPT)", ok if tem_prints else pend, "docs/Documentacao_DevOps.pptx"],
        ["Deploy realizado nos ambientes staging e produção", ok if DEPLOY_OK else pend, "Render + Supabase"],
    ], size=14, row_h=0.6)

    # 17. Links
    s = d.dark()
    text(s, 0.8, 1.2, 11.5, 0.4, "ONDE VER FUNCIONANDO", size=14, bold=True, color=ACCENT)
    text(s, 0.8, 1.65, 11.5, 0.9, "Versionado, testado e publicado automaticamente", size=32, bold=True, color="FFFFFF")
    links = [("Repositório", REPO_URL), ("Pipeline", REPO_URL + "/actions"),
             ("Staging", STAGING_URL + "/swagger"), ("Produção", PROD_URL + "/swagger")]
    for i, (label, url) in enumerate(links):
        y = 3.05 + i * 0.85
        box(s, 0.8, y, 11.7, 0.68, fill=DARK2).name = f"Link {label}"
        text(s, 1.1, y + 0.12, 2.2, 0.45, label, size=16, bold=True, color=ACCENT, anchor=MSO_ANCHOR.MIDDLE)
        text(s, 3.3, y + 0.12, 9.0, 0.45, url, size=16, color="FFFFFF", anchor=MSO_ANCHOR.MIDDLE, font=CODE_FONT)

    d.footers()
    return d.prs


if __name__ == "__main__":
    prs = build()
    SAIDA.parent.mkdir(parents=True, exist_ok=True)
    prs.save(SAIDA)
    print(f"Gerado {SAIDA.relative_to(ROOT)} ({len(prs.slides)} slides)")
